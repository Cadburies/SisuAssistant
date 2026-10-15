#!/usr/bin/env python3
"""Build the PCBWay fab + assembly package into production/ from the KiCad project.

    python3 fab_package.py            # KiCad must be closed (reads the saved files)

Wipes production/ (except the Quilter net-class CSV, a routing input), then writes:
  gerbers/ + MarineBoard-PCBWay-Gerbers.zip   4 copper, mask, paste, silk, outline,
                                               PTH/NPTH Excellon + maps, IPC-D-356 netlist
  MarineBoard-BOM.csv / .xlsx                 grouped BOM + plug-in fuses and the JP1 shunt
  MarineBoard-positions.csv / .xlsx           centroid (xlsx = PCBWay upload)

Refills zones first; aborts if DRC (errors only, with schematic parity) is not clean.
"""
import csv, os, re, shutil, subprocess, sys, zipfile
import openpyxl
from openpyxl.styles import Font

KICAD_CLI = os.environ.get("KICAD_CLI", "/Applications/KiCad/KiCad.app/Contents/MacOS/kicad-cli")
HERE = os.path.dirname(os.path.abspath(__file__))
PCB = os.path.join(HERE, "MarineBoard.kicad_pcb")
SCH = os.path.join(HERE, "MarineBoard.kicad_sch")
OUT = os.path.join(HERE, "production")
GERB = os.path.join(OUT, "gerbers")
KEEP = {"quilter_high_current_nets.csv"}
LAYERS = "F.Cu,In1.Cu,In2.Cu,B.Cu,F.Paste,B.Paste,F.SilkS,B.SilkS,F.Mask,B.Mask,Edge.Cuts"

# Plug-in parts. The holder or header is soldered. These go in afterwards.
# Littelfuse MINI 297 starts at 2 A, so F2 is Eaton Bussmann ATM-1 (1 A, same MINI body).
# refs, value, mpn, description, manufacturer, qty, instruction
PLUG_IN = [
    ("F1-FUSE", "2A", "0297002.WXNV",
     "MINI blade fuse 2 A 32 V. Insert into holder F1.", "Littelfuse", 1,
     "Insert into holder F1"),
    ("F2-FUSE", "1A", "ATM-1",
     "MINI blade fuse 1 A 32 V. Insert into holder F2.", "Eaton Bussmann", 1,
     "Insert into holder F2"),
    ("F3-FUSE,F4-FUSE", "10A", "0297010.WXNV",
     "MINI blade fuse 10 A 32 V. Insert one into holder F3 and one into holder F4.", "Littelfuse", 2,
     "Insert into holders F3 and F4"),
    ("JP1-SHUNT", "1.27mm 1x2", "1.27",
     "Closed-top 1x2 jumper shunt, 1.27 mm pitch, LCSC C150949. Fit onto header JP1.", "BOOMELE", 1,
     "Fit onto header JP1"),
]
INSTRUCTIONS = {"JP1": "Fit the BOOMELE 1.27 shunt, LCSC C150949, onto these pins"}


def cli(*args):
    r = subprocess.run([KICAD_CLI, *args], capture_output=True, text=True)
    if r.returncode:
        sys.exit(f"kicad-cli {' '.join(args[:3])} failed:\n{r.stdout}{r.stderr}")
    return r.stdout


def drc_gate():
    rpt = os.path.join(OUT, ".drc.rpt")
    subprocess.run([KICAD_CLI, "pcb", "drc", "--schematic-parity", "--refill-zones", "--save-board",
                    "--severity-error", "-o", rpt, PCB], capture_output=True)
    text = open(rpt).read()
    os.remove(rpt)
    counts = [int(n) for n in re.findall(r"\*\* Found (\d+)", text)]
    if any(counts):
        sys.exit("DRC errors - fix before building the fab package:\n" + text)


def footprint_info():
    """ref -> (side, 'SMD'|'THT') from the saved board."""
    t = open(PCB).read()
    info = {}
    for m in re.finditer(r'\n\t\(footprint "[^"]+"\n(.*?)\n\t\)', t, re.S):
        body = m.group(1)
        ref = re.search(r'\(property "Reference" "([^"]+)"', body).group(1)
        side = "Top" if re.search(r'\(layer "([^"]+)"\)', body).group(1) == "F.Cu" else "Bottom"
        attr = re.search(r"\(attr ([^)]*)\)", body)
        kind = "THT" if attr and "through_hole" in attr.group(1) else "SMD"
        info[ref] = (side, kind)
    return info


def fmt_sheet(ws, widths):
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def main():
    os.makedirs(OUT, exist_ok=True)
    for name in os.listdir(OUT):
        if name in KEEP:
            continue
        p = os.path.join(OUT, name)
        shutil.rmtree(p) if os.path.isdir(p) else os.remove(p)
    drc_gate()

    os.makedirs(GERB)
    cli("pcb", "export", "gerbers", "-l", LAYERS, "-o", GERB + "/", PCB)
    cli("pcb", "export", "drill", "--format", "excellon", "--excellon-separate-th", "--excellon-units", "mm",
        "--generate-map", "--map-format", "gerberx2", "-o", GERB + "/", PCB)
    cli("pcb", "export", "ipcd356", "-o", os.path.join(GERB, "MarineBoard-netlist.ipc"), PCB)
    with zipfile.ZipFile(os.path.join(OUT, "MarineBoard-PCBWay-Gerbers.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(os.listdir(GERB)):
            z.write(os.path.join(GERB, f), f)

    # BOM
    bom_csv = os.path.join(OUT, "MarineBoard-BOM.csv")
    cli("sch", "export", "bom", "--exclude-dnp",
        "--fields", "Reference,Value,Footprint,${QUANTITY},MPN,Description,Manufacturer",
        "--labels", "Refs,Value,Footprint,Qty,MPN,Description,Manufacturer",
        "--group-by", "Value,Footprint,MPN", "--sort-field", "Reference", "--ref-range-delimiter", "",
        "-o", bom_csv, SCH)
    with open(bom_csv, "a", newline="") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL)
        for refs, val, mpn, desc, mfr, qty, _instr in PLUG_IN:
            w.writerow([refs, val, "INSERT", qty, mpn, desc, mfr])

    fps = footprint_info()
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "BOM"
    ws.append(["Item #", "Designator", "Qty", "Manufacturer", "Mfg Part # / LCSC", "Description", "Value",
               "Package", "Type", "Side", "Instructions"])
    for i, r in enumerate(csv.DictReader(open(bom_csv)), 1):
        refs = r["Refs"].split(",")
        if r["Footprint"] == "INSERT":
            instr = next(row[6] for row in PLUG_IN if row[0] == r["Refs"])
            ws.append([i, r["Refs"], int(r["Qty"]), r["Manufacturer"] or None, r["MPN"] or None, r["Description"],
                       r["Value"], "MINI blade" if r["Refs"].startswith("F") else "1.27 mm shunt",
                       "Plug-in", "Top", instr])
            continue
        missing = [x for x in refs if x not in fps]
        if missing:
            sys.exit(f"BOM refs not on the board: {missing}")
        sides = sorted({fps[x][0] for x in refs})
        kinds = {fps[x][1] for x in refs}
        pkg = r["Footprint"].split(":")[-1]
        if pkg == "D_SOD-882":
            pkg = "DFN1006-2L"
        ws.append([i, r["Refs"], int(r["Qty"]), r["Manufacturer"] or None, r["MPN"] or None, r["Description"] or None,
                   r["Value"], pkg, "THT" if kinds == {"THT"} else "SMD", "/".join(sides),
                   next((INSTRUCTIONS[x] for x in refs if x in INSTRUCTIONS), None)])
    fmt_sheet(ws, {"A": 7, "B": 26, "C": 5, "D": 16, "E": 20, "F": 60, "G": 22, "H": 34, "I": 7, "J": 11, "K": 30})
    wb.save(os.path.join(OUT, "MarineBoard-BOM.xlsx"))

    # Centroid
    pos_csv = os.path.join(OUT, "MarineBoard-positions.csv")
    cli("pcb", "export", "pos", "--format", "csv", "--units", "mm", "--side", "both", "--exclude-dnp", "-o", pos_csv, PCB)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Centroid"
    ws.append(["Designator", "Value", "Package", "Mid X (mm)", "Mid Y (mm)", "Rotation", "Layer"])
    for r in csv.DictReader(open(pos_csv)):
        ws.append([r["Ref"], r["Val"], r["Package"], float(r["PosX"]), float(r["PosY"]), float(r["Rot"]),
                   r["Side"].capitalize()])
    fmt_sheet(ws, {"A": 12, "B": 24, "C": 34, "D": 12, "E": 12, "F": 10, "G": 8})
    wb.save(os.path.join(OUT, "MarineBoard-positions.xlsx"))

    print("production/:", ", ".join(sorted(os.listdir(OUT))))
    print("gerbers/:", ", ".join(sorted(os.listdir(GERB))))


if __name__ == "__main__":
    main()
